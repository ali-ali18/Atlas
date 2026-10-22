"""Testes locais do doclib/build_reference (sem LibreOffice/pandoc/TeX; o que precisa deles e do smoke.py).

    python -m unittest docker/code-runner/test_doclib.py   (precisa de python-docx; sem ele os testes sao pulados)
"""
import os
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
os.environ.setdefault("DOCLIB_TEMPLATES", os.path.join(HERE, "templates", "abnt"))

import doclib  # noqa: E402

try:
    import docx  # noqa: F401
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


class FillTemplate(unittest.TestCase):
    def test_substitui_marcadores_e_copia_bib(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "t.tex")
            doclib.fill_template("abnt-trabalho", out, titulo="Meu título", AUTOR="Ana", TEXTO=r"\chapter{Introdução} Olá.")
            text = open(out, encoding="utf-8").read()
            self.assertIn(r"\titulo{Meu título}", text)
            self.assertIn(r"\chapter{Introdução} Olá.", text)
            self.assertNotIn("<<", text)
            self.assertTrue(os.path.exists(os.path.join(tmp, "referencias.bib")))

    def test_projeto_tem_secoes_padrao(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "p.tex")
            doclib.fill_template("abnt-projeto", out, TITULO="X")
            text = open(out, encoding="utf-8").read()
            for name in ("Metodologia", "Cronograma", "Justificativa"):
                self.assertIn(name, text)


@unittest.skipUnless(HAS_DOCX, "python-docx ausente")
class Docx(unittest.TestCase):
    def test_referencia_abnt_tem_estilos_e_pagina(self):
        import build_reference
        from docx import Document
        from docx.shared import Cm

        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "ref.docx")
            build_reference.build(out)
            doc = Document(out)
            names = {s.name for s in doc.styles}
            for needed in ("Body Text", "First Paragraph", "Block Text", "Table Caption", "Image Caption", "Fonte", "Bibliography", "Compact", "toc 1", "Heading 1"):
                self.assertIn(needed, names)
            sec = doc.sections[0]
            self.assertAlmostEqual(sec.left_margin.cm, 3.0, places=1)
            self.assertAlmostEqual(sec.bottom_margin.cm, 2.0, places=1)
            self.assertAlmostEqual(doc.styles["Body Text"].paragraph_format.first_line_indent.cm, 1.25, places=2)
            self.assertAlmostEqual(doc.styles["Block Text"].paragraph_format.left_indent.cm, 4.0, places=1)
            self.assertEqual(doc.styles["Block Text"].font.size.pt, 10)
            self.assertTrue(doc.styles["Heading 1"].font.all_caps)
            self.assertIn("PAGE", sec.header.paragraphs[0]._p.xml)
            _ = Cm

    def test_docx_add_toc_antes_do_primeiro_titulo(self):
        import build_reference
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "d.docx")
            build_reference.build(path)
            doc = Document(path)
            doc.add_paragraph("Capa")
            doc.add_heading("Introdução", 1)
            doc.save(path)
            doclib.docx_add_toc(path)
            texts = [p.text for p in Document(path).paragraphs]
            self.assertLess(texts.index("SUMÁRIO"), texts.index("Introdução"))
            self.assertLess(texts.index("Capa"), texts.index("SUMÁRIO"))
            xml = Document(path).element.xml
            self.assertIn("TOC", xml)

    def test_docx_add_toc_sem_titulos_falha(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "d.docx")
            Document().save(path)
            with self.assertRaises(ValueError):
                doclib.docx_add_toc(path)


if __name__ == "__main__":
    unittest.main()
