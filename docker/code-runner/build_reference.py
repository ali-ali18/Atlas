"""Gera o documento de referencia ABNT para o pandoc (`pandoc --reference-doc`).

Rodado no build da imagem (`python build_reference.py /app/templates/abnt/abnt-reference.docx`);
nada de binario no repositorio. Estilos: A4, margens 3/2 cm, Times New Roman 12 (a imagem renderiza
com Liberation Serif, metrica identica), entrelinha 1,5, recuo de 1,25 cm, justificado.
Para estilos proprios use blocos do pandoc: `::: {custom-style="Fonte"}` ... `:::`.
"""
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
J, L, C = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER


def style(doc, name, kind=WD_STYLE_TYPE.PARAGRAPH, base="Normal"):
    styles = doc.styles
    try:
        st = styles[name]
    except KeyError:
        st = styles.add_style(name, kind)
        if base and kind == WD_STYLE_TYPE.PARAGRAPH:
            st.base_style = styles[base]
    return st


def font(st, size=12, bold=False, italic=False, caps=False):
    f = st.font
    f.name, f.size, f.bold, f.italic, f.all_caps = FONT, Pt(size), bold, italic, caps
    f.color.rgb = RGBColor(0, 0, 0)
    rpr = st.element.get_or_add_rPr()
    for el in rpr.findall(qn("w:color")):
        for attr in (qn("w:themeColor"), qn("w:themeShade"), qn("w:themeTint")):
            el.attrib.pop(attr, None)
    fonts = rpr.find(qn("w:rFonts"))
    for attr in list(fonts.attrib):
        if "Theme" in attr:
            del fonts.attrib[attr]
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn("w:" + attr), FONT)


def para(st, align=J, spacing=1.5, indent=Cm(0), left=Cm(0), before=0, after=0, keep_next=False):
    p = st.paragraph_format
    p.alignment, p.line_spacing = align, spacing
    p.first_line_indent, p.left_indent, p.right_indent = indent, left, Cm(0)
    p.space_before, p.space_after = Pt(before), Pt(after)
    p.keep_with_next = keep_next
    p.widow_control = True


def page_field(paragraph, instr):
    """Campo complexo (begin/instr/separate/texto/end), um run por parte como o Word grava."""
    for kind in ("begin", "instr", "separate", "text", "end"):
        run = paragraph.add_run()
        run.font.name, run.font.size = FONT, Pt(10)
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        elif kind == "text":
            el = OxmlElement("w:t")
            el.text = "1"
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def build(out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.left_margin, sec.bottom_margin, sec.right_margin = Cm(3), Cm(3), Cm(2), Cm(2)
    sec.header_distance, sec.footer_distance = Cm(1.25), Cm(1.25)

    # Corpo: pandoc usa "Body Text" e "First Paragraph" (ABNT recua tambem o primeiro paragrafo).
    normal = doc.styles["Normal"]
    font(normal)
    para(normal, indent=Cm(1.25))
    for name in ("Body Text", "First Paragraph"):
        st = style(doc, name)
        font(st)
        para(st, indent=Cm(1.25))
    st = style(doc, "Compact")  # itens de lista e celulas de tabela
    font(st)
    para(st, align=L, spacing=1.0)

    # Titulos: 1 caixa alta negrito; 2 negrito; 3 negrito italico; 4+ italico.
    levels = [(True, False, True), (True, False, False), (True, True, False), (False, True, False), (False, True, False)]
    for level, (bold, italic, caps) in enumerate(levels, 1):
        st = doc.styles[f"Heading {level}"]
        font(st, 12, bold, italic, caps)
        para(st, align=L, spacing=1.5, before=18, after=12, keep_next=True)
    for name in ("Title", "Subtitle", "Author", "Date"):
        st = style(doc, name)
        font(st, 14 if name == "Title" else 12, bold=name == "Title", caps=name == "Title")
        para(st, align=C, spacing=1.5)
    st = style(doc, "Abstract")  # resumo (metadado `abstract`)
    font(st)
    para(st, spacing=1.0)

    # Citacao longa (> 3 linhas): recuo 4 cm, 10 pt, espacamento simples.
    st = style(doc, "Block Text")
    font(st, 10)
    para(st, spacing=1.0, left=Cm(4), before=6, after=12)

    # Legendas e fonte: 10 pt, simples.
    for name in ("Caption", "Image Caption", "Table Caption", "Figure", "Captioned Figure"):
        st = style(doc, name)
        font(st, 10)
        para(st, align=C, spacing=1.0, before=6, after=6, keep_next=name in ("Table Caption", "Caption"))
    st = style(doc, "Fonte")
    font(st, 10)
    para(st, align=L, spacing=1.0, after=12)

    # Notas de rodape: 10 pt, simples.
    st = style(doc, "Footnote Text")
    font(st, 10)
    para(st, spacing=1.0)

    # Referencias: simples, a esquerda, espaco depois (estilo Bibliography ou custom-style="Referencias").
    for name in ("Bibliography", "Referencias"):
        st = style(doc, name)
        font(st)
        para(st, align=L, spacing=1.0, after=12)

    # Sumario.
    st = style(doc, "TOC Heading")
    font(st, 12, bold=True, caps=True)
    para(st, align=C, spacing=1.5, after=12)
    for level, bold in ((1, True), (2, False), (3, False)):
        st = style(doc, f"toc {level}")
        font(st, 12, bold=bold, caps=level == 1)
        para(st, align=L, spacing=1.5, left=Cm(0.75 * (level - 1)))
        st.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    # Codigo e hiperlink.
    st = style(doc, "Source Code")
    st.font.name, st.font.size = "Liberation Mono", Pt(10)
    para(st, align=L, spacing=1.0)
    st = style(doc, "Hyperlink", WD_STYLE_TYPE.CHARACTER)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.underline = True

    # Tabela no estilo IBGE: so linhas horizontais (topo, cabecalho e base), sem bordas laterais.
    tbl = parse_xml(
        f'<w:style {nsdecls("w")} w:type="table" w:customStyle="1" w:styleId="Table"><w:name w:val="Table"/>'
        '<w:basedOn w:val="TableNormal"/><w:uiPriority w:val="59"/><w:tblPr><w:tblInd w:w="0" w:type="dxa"/>'
        '<w:tblBorders><w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/></w:tblBorders>'
        '<w:tblCellMar><w:top w:w="40" w:type="dxa"/><w:left w:w="100" w:type="dxa"/>'
        '<w:bottom w:w="40" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tblCellMar></w:tblPr>'
        '<w:tblStylePr w:type="firstRow"><w:rPr><w:b/></w:rPr><w:tcPr><w:tcBorders>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/></w:tcBorders></w:tcPr></w:tblStylePr></w:style>'
    )
    for old in doc.styles.element.findall(qn("w:style")):
        if old.get(qn("w:styleId")) == "Table":
            doc.styles.element.remove(old)
    doc.styles.element.append(tbl)

    # Cabecalho: numero da pagina alinhado a direita.
    header = sec.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.style = doc.styles["Header"]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.first_line_indent = Cm(0)
    hp.paragraph_format.line_spacing = 1.0
    page_field(hp, "PAGE")

    doc.save(out)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "abnt-reference.docx")
