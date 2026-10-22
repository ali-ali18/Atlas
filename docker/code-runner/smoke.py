"""Teste de fumaca da imagem code-runner (rodar DENTRO da imagem, como no runtime):

    docker run --rm --network none --read-only --user 10001:10001 \
      --tmpfs /tmp:size=128m,mode=1777 --tmpfs /work:size=160m,uid=10001,gid=10001 \
      atlas-code-runner:latest python /app/smoke.py

Sai com codigo != 0 se algum passo falhar; imprime OK/FALHA por passo.
"""
import os
import shutil
import subprocess
import sys
import tempfile
import traceback

sys.path.insert(0, "/app")
import doclib  # noqa: E402

WORK = "/work/smoke"
MD = """---
title: Teste
author: Autora
---

# Introdução

Primeiro parágrafo do trabalho, com texto suficiente para ocupar linhas e verificar recuo e justificação.

## Citação longa

> Citação com mais de três linhas, recuada em quatro centímetros, em fonte menor e espaçamento simples, para conferir o estilo.

Tabela 1 – Exemplo

| Item | Valor |
|------|------:|
| A    | 1     |
| B    | 2     |

::: {custom-style="Fonte"}
Fonte: elaborado pelo autor (2025).
:::

# Conclusão

Fim.
"""
TEX = r"""\documentclass[12pt, oneside, a4paper, chapter=TITLE, section=TITLE, sumario=abnt-6027-2012, english, brazil]{abntex2}
\usepackage[T1]{fontenc}
\usepackage[utf8]{inputenc}
\usepackage[alf]{abntex2cite}
\titulo{Teste}\autor{Autora}\local{Cidade}\data{2025}\orientador{Prof.}\instituicao{Instituição}
\tipotrabalho{Trabalho}\preambulo{Preâmbulo}
\begin{document}
\imprimircapa
\imprimirfolhaderosto*
\tableofcontents*
\textual
\chapter{Introdução}
Texto com citação \cite{exemplo2020}.
\postextual
\bibliography{referencias}
\end{document}
"""


def sh(*cmd):
    return subprocess.run(cmd, capture_output=True, text=True, timeout=110, cwd=WORK)


def step_python_libs():
    import docx, pptx, pypdf, pdfplumber, openpyxl, pandas, numpy, scipy, matplotlib  # noqa: F401, E401
    from docx import Document
    d = Document()
    d.add_heading("Olá", 1)
    d.add_paragraph("corpo")
    d.save(f"{WORK}/py.docx")
    assert Document(f"{WORK}/py.docx").paragraphs[0].text == "Olá"


def step_pandoc():
    with open(f"{WORK}/t.md", "w", encoding="utf-8") as fh:
        fh.write(MD)
    doclib.md_to_docx_abnt(f"{WORK}/t.md", f"{WORK}/t.docx", toc=True)
    assert os.path.getsize(f"{WORK}/t.docx") > 5000


def step_soffice():
    pdf = doclib.docx_to_pdf(f"{WORK}/t.docx", outdir=WORK)
    import pdfplumber
    with pdfplumber.open(pdf) as doc:
        text = "\n".join(page.extract_text() or "" for page in doc.pages)
    assert "INTRODUÇÃO" in text.upper(), text[:300]
    # sumario atualizado pelo LibreOffice deve listar a Conclusao
    assert text.upper().count("CONCLUSÃO") >= 2, "sumario nao foi atualizado:\n" + text[:500]


def step_pdftoppm():
    pngs = doclib.pdf_to_pngs(f"{WORK}/t.pdf", dpi=60, outdir=f"{WORK}/preview")
    assert pngs and all(os.path.getsize(p) > 1000 for p in pngs), pngs


def step_toc_helper():
    from docx import Document
    d = Document(f"{WORK}/py.docx")
    d.save(f"{WORK}/py2.docx")
    doclib.docx_add_toc(f"{WORK}/py2.docx")


def step_latex():
    with open(f"{WORK}/a.tex", "w", encoding="utf-8") as fh:
        fh.write(TEX)
    shutil.copy("/app/templates/abnt/referencias.bib", f"{WORK}/referencias.bib")
    pdf = doclib.latex_abnt(f"{WORK}/a.tex")
    import pdfplumber
    with pdfplumber.open(pdf) as doc:
        assert len(doc.pages) >= 3
    doclib.fill_template("abnt-trabalho", f"{WORK}/b.tex", TITULO="T", AUTOR="A", TEXTO=r"\chapter{Intro} Texto \cite{exemplo2020}.", RESUMO="r", ABSTRACT="a", PALAVRAS_CHAVE="p", KEYWORDS="k")
    doclib.latex_abnt(f"{WORK}/b.tex")


STEPS = [step_python_libs, step_pandoc, step_soffice, step_pdftoppm, step_toc_helper, step_latex]


def main():
    shutil.rmtree(WORK, ignore_errors=True)
    os.makedirs(WORK)
    failed = 0
    for step in STEPS:
        try:
            step()
            print("OK    ", step.__name__)
        except Exception:
            failed += 1
            print("FALHA ", step.__name__)
            traceback.print_exc()
    print("uid", os.getuid(), "HOME", os.environ.get("HOME"), "tmp", tempfile.gettempdir())
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
