"""Funcoes de documento para o `run_code` (imagem code-runner). Uso:

    import sys; sys.path.insert(0, "/app")
    import doclib
    doclib.md_to_docx_abnt("/work/in/trabalho.md", "/work/out/trabalho.docx")
    pdf = doclib.docx_to_pdf("/work/out/trabalho.docx", outdir="/work")      # PDF fora de /work/out
    pages = doclib.pdf_to_pngs(pdf)                                          # PNGs em /work/preview (voltam como imagens)

PNGs em /work/preview sao devolvidos ao modelo (ate 12) e NAO sao entregues ao usuario; o que vai
para o usuario fica em /work/out/. Estilos para blocos do pandoc (`::: {custom-style="Fonte"}`):
Fonte, Referencias, Block Text (citacao longa), Table Caption, Image Caption.
"""
from __future__ import annotations

import glob
import os
import re
import shutil
import subprocess

TEMPLATES = os.environ.get("DOCLIB_TEMPLATES", "/app/templates/abnt")
LO_PROFILE = "-env:UserInstallation=file:///tmp/lo"
PREVIEW_DIR = "/work/preview"


def _run(cmd: list[str], timeout: int = 100, cwd: str | None = None) -> subprocess.CompletedProcess:
    """Executa um comando; falha com a saida do programa (o modelo precisa ler o erro)."""
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, cwd=cwd)
    if proc.returncode != 0:
        tail = (proc.stdout + "\n" + proc.stderr).strip()[-3000:]
        raise RuntimeError(f"{cmd[0]} falhou (codigo {proc.returncode}):\n{tail}")
    return proc


def md_to_docx_abnt(md_path: str, out_path: str, toc: bool = False, number_sections: bool = True, extra_args: list[str] | None = None) -> str:
    """Markdown -> DOCX com os estilos ABNT (pandoc + abnt-reference.docx). Titulos numerados (1, 1.1...).

    toc=True insere um campo de sumario (preenchido pelo Word ao abrir e pelo docx_to_pdf).
    """
    cmd = ["pandoc", md_path, "-f", "markdown+smart", "-o", out_path, "--reference-doc", os.path.join(TEMPLATES, "abnt-reference.docx"), "-M", "lang=pt-BR", "-M", "toc-title=SUMÁRIO"]
    if number_sections:
        cmd.append("--number-sections")
    if toc:
        cmd += ["--toc", "--toc-depth=3"]
    _run(cmd + list(extra_args or []), cwd=os.path.dirname(os.path.abspath(md_path)))
    return out_path


def docx_to_pdf(path: str, outdir: str | None = None, update_fields: bool = True) -> str:
    """DOCX (ou ODT/PPTX/XLSX) -> PDF com o LibreOffice; retorna o caminho do PDF.

    Com update_fields (padrao) o sumario e os campos sao atualizados antes de exportar. Salve o PDF de
    conferencia fora de /work/out (ex.: outdir="/work") para nao entrega-lo sem querer.
    """
    outdir = outdir or os.path.dirname(os.path.abspath(path))
    os.makedirs(outdir, exist_ok=True)
    pdf = os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + ".pdf")
    if update_fields and path.lower().endswith(".docx") and os.path.exists("/usr/bin/python3") and os.path.exists("/app/lo_export.py"):
        try:
            _run(["/usr/bin/python3", "/app/lo_export.py", path, pdf], timeout=110)
            if os.path.exists(pdf):
                return pdf
        except Exception:
            pass  # cai na conversao simples (sem atualizar o sumario)
    _run(["soffice", "--headless", "--norestore", LO_PROFILE, "--convert-to", "pdf", "--outdir", outdir, path], timeout=110)
    if not os.path.exists(pdf):
        raise RuntimeError("LibreOffice nao gerou o PDF")
    return pdf


def pdf_to_pngs(pdf: str, dpi: int = 60, outdir: str = PREVIEW_DIR, first: int = 1, max_pages: int = 12) -> list[str]:
    """Rasteriza paginas do PDF (poppler) em PNGs `p-NN.png`; retorna os caminhos em ordem.

    So as `max_pages` paginas a partir de `first` (o runtime devolve no maximo 12 imagens por chamada);
    use dpi=60 para ver o layout e dpi=90-110 para conferir texto miudo.
    """
    os.makedirs(outdir, exist_ok=True)
    for old in glob.glob(os.path.join(outdir, "p-*.png")):
        os.remove(old)
    _run(["pdftoppm", "-r", str(dpi), "-png", "-f", str(first), "-l", str(first + max_pages - 1), pdf, os.path.join(outdir, "p")])
    return sorted(glob.glob(os.path.join(outdir, "p-*.png")), key=lambda name: int(re.findall(r"(\d+)\.png$", name)[0]))


def fill_template(name: str, out_path: str, **fields: str) -> str:
    """Copia um modelo .tex (abnt-trabalho / abnt-projeto) trocando <<CHAVE>> por fields[CHAVE].

    Ex.: fill_template("abnt-trabalho", "/work/t.tex", TITULO="...", AUTOR="...", TEXTO=r"\\chapter{Introducao} ...").
    Tambem grava referencias.bib (modelo) ao lado se ainda nao existir. Marcadores sem valor ficam vazios.
    """
    src = os.path.join(TEMPLATES, name if name.endswith(".tex") else name + ".tex")
    with open(src, encoding="utf-8") as fh:
        text = fh.read()
    values = {key.upper(): str(value) for key, value in fields.items()}
    text = re.sub(r"<<([A-Z_]+)>>", lambda m: values.get(m.group(1), ""), text)
    with open(out_path, "w", encoding="utf-8") as fh:
        fh.write(text)
    bib = os.path.join(os.path.dirname(os.path.abspath(out_path)), "referencias.bib")
    if not os.path.exists(bib):
        shutil.copy(os.path.join(TEMPLATES, "referencias.bib"), bib)
    return out_path


def latex_abnt(tex_path: str) -> str:
    """Compila um .tex abnTeX2 com latexmk (pdflatex + bibtex, offline); retorna o PDF ao lado do .tex."""
    tex_path = os.path.abspath(tex_path)
    folder = os.path.dirname(tex_path)
    try:
        _run(["latexmk", "-r", os.path.join(TEMPLATES, "latexmkrc"), os.path.basename(tex_path)], timeout=110, cwd=folder)
    except RuntimeError as error:
        log = tex_path[:-4] + ".log"
        hint = ""
        if os.path.exists(log):
            with open(log, encoding="utf-8", errors="replace") as fh:
                lines = [line for line in fh.read().splitlines() if line.startswith("!") or "Error" in line]
            hint = "\nLinhas de erro do .log:\n" + "\n".join(lines[:15])
        raise RuntimeError(str(error) + hint) from None
    pdf = tex_path[:-4] + ".pdf"
    if not os.path.exists(pdf):
        raise RuntimeError("latexmk nao gerou o PDF")
    return pdf


def docx_add_toc(path: str, title: str = "SUMÁRIO", levels: int = 3) -> str:
    """Insere um sumario (campo TOC dos titulos 1..levels) no DOCX, antes do primeiro 'Heading 1'.

    Se existir um paragrafo com o texto exato [SUMARIO] o sumario entra ali. O campo e marcado para o Word
    atualizar ao abrir; docx_to_pdf ja o atualiza na conferencia. Retorna o mesmo caminho.
    """
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document(path)
    anchor = next((p for p in doc.paragraphs if p.text.strip() == "[SUMARIO]"), None)
    replace = anchor is not None
    if anchor is None:
        anchor = next((p for p in doc.paragraphs if p.style.name == "Heading 1"), None)
    if anchor is None:
        raise ValueError("nenhum titulo (Heading 1) nem marcador [SUMARIO] no documento")
    style_names = {s.name for s in doc.styles}

    def fld(par, kind=None, instr=None, text=None):
        run = par.add_run()
        if instr is not None:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        elif text is not None:
            el = OxmlElement("w:t")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            if kind == "begin":
                el.set(qn("w:dirty"), "true")
        run._r.append(el)

    heading = anchor.insert_paragraph_before(title)
    heading.style = doc.styles["TOC Heading" if "TOC Heading" in style_names else "Normal"]
    body = anchor.insert_paragraph_before()
    fld(body, "begin")
    fld(body, instr=f' TOC \\o "1-{levels}" \\h \\z \\u ')
    fld(body, "separate")
    fld(body, text="Clique com o botão direito e escolha Atualizar campo para gerar o sumário.")
    fld(body, "end")
    brk = anchor.insert_paragraph_before()
    brk.add_run().add_break(WD_BREAK.PAGE)
    if replace:
        anchor._p.getparent().remove(anchor._p)
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        settings.append(upd)
    doc.save(path)
    return path
